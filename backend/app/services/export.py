"""Meeting minutes as a Word (.docx) document, in the meeting's minutes language (Arabic = right-to-left).

PDF export is done by the browser (print the minutes page), which shapes Arabic correctly for free.
"""

from datetime import datetime
from io import BytesIO

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

from ..config import settings
from ..models import Meeting

FONT = "Arial"  # ships with Word everywhere and has full Arabic glyphs

LABELS = {
    "ar": {
        "minutes": "محضر اجتماع",
        "date": "التاريخ",
        "time": "الوقت",
        "duration": "المدة",
        "attendees": "الحضور",
        "chair": "رئيس الجلسة",
        "agenda": "جدول الأعمال",
        "covered": "تمت مناقشته",
        "skipped": "لم يُناقش",
        "summary": "الملخص",
        "decisions": "القرارات",
        "tasks": "المهام",
        "task": "المهمة",
        "owner": "المسؤول",
        "due": "الموعد",
        "status": "الحالة",
        "done": "منجزة",
        "open": "مفتوحة",
        "transcript": "النص الكامل",
        "slide": "الشريحة",
        "draft": "مسودة — لم يُعتمد بعد",
        "approved": "معتمد",
        "minutesUnit": "دقيقة",
        "none": "لا يوجد.",
    },
    "en": {
        "minutes": "Meeting minutes",
        "date": "Date",
        "time": "Time",
        "duration": "Duration",
        "attendees": "Attendees",
        "chair": "Chair",
        "agenda": "Agenda",
        "covered": "Discussed",
        "skipped": "Not discussed",
        "summary": "Summary",
        "decisions": "Decisions",
        "tasks": "Tasks",
        "task": "Task",
        "owner": "Owner",
        "due": "Due",
        "status": "Status",
        "done": "Done",
        "open": "Open",
        "transcript": "Full transcript",
        "slide": "slide",
        "draft": "Draft — not approved yet",
        "approved": "Approved",
        "minutesUnit": "min",
        "none": "None.",
    },
}


def _rtl_paragraph(paragraph) -> None:
    # Right-to-left paragraph. No explicit alignment: an RTL paragraph already starts on the right,
    # and jc="right" inside one is read as the *far* side (left) by Word and LibreOffice.
    paragraph._p.get_or_add_pPr().insert(0, OxmlElement("w:bidi"))


def _style_run(run, rtl: bool, size: float = 11, bold: bool = False, color: RGBColor | None = None) -> None:
    run.font.name = FONT
    run.font.size = Pt(size)
    run.bold = bold
    if color:
        run.font.color.rgb = color
    rpr = run._r.get_or_add_rPr()
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.append(fonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs"):  # cs = the font Word uses for Arabic
        fonts.set(qn(attr), FONT)
    size_cs = OxmlElement("w:szCs")
    size_cs.set(qn("w:val"), str(int(size * 2)))
    rpr.append(size_cs)
    if bold:
        rpr.append(OxmlElement("w:bCs"))
    if rtl:
        rpr.append(OxmlElement("w:rtl"))


class _Writer:
    def __init__(self, rtl: bool) -> None:
        self.rtl = rtl
        self.doc = Document()
        for section in self.doc.sections:
            if rtl:
                section._sectPr.append(OxmlElement("w:bidi"))

    def para(self, text: str = "", size: float = 11, bold: bool = False, color: RGBColor | None = None, space_after: float = 4):
        p = self.doc.add_paragraph()
        if self.rtl:
            _rtl_paragraph(p)
        p.paragraph_format.space_after = Pt(space_after)
        if text:
            _style_run(p.add_run(text), self.rtl, size, bold, color)
        return p

    def heading(self, text: str) -> None:
        p = self.para(text, size=13, bold=True, space_after=6)
        p.paragraph_format.space_before = Pt(14)

    def table(self, header: list[str] | None, rows: list[list[str]]) -> None:
        table = self.doc.add_table(rows=1 if header else 0, cols=len(header or rows[0]))
        table.style = "Table Grid"
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        if self.rtl:
            table._tbl.tblPr.append(OxmlElement("w:bidiVisual"))  # columns run right-to-left
        if header:
            for cell, text in zip(table.rows[0].cells, header):
                self._cell(cell, text, bold=True)
        for row in rows:
            cells = table.add_row().cells
            for cell, text in zip(cells, row):
                self._cell(cell, text)

    def _cell(self, cell, text: str, bold: bool = False) -> None:
        p = cell.paragraphs[0]
        if self.rtl:
            _rtl_paragraph(p)
        _style_run(p.add_run(text), self.rtl, 10, bold)


def _hm(dt: datetime | None) -> str:
    return dt.strftime("%H:%M") if dt else "—"


def minutes_docx(meeting: Meeting) -> bytes:
    lang = meeting.language.value if meeting.language.value in LABELS else "ar"
    L = LABELS[lang]
    w = _Writer(rtl=lang == "ar")
    muted = RGBColor(0x6A, 0x6A, 0x66)
    names = {p.id: p.name for p in meeting.participants}
    report = meeting.report

    # Letterhead
    if settings.org_name:
        w.para(settings.org_name, size=12, bold=True, color=muted, space_after=0)
    w.para(L["minutes"], size=10, color=muted, space_after=2)
    w.para(meeting.title, size=20, bold=True, space_after=2)
    if report:
        w.para(L["approved"] if report.approved else L["draft"], size=9, color=muted, space_after=10)

    # Details
    minutes = (
        round((meeting.ended_at - meeting.started_at).total_seconds() / 60) if meeting.started_at and meeting.ended_at else None
    )
    chair = next((p.name for p in meeting.participants if p.is_chair), None)
    attendees = "، ".join(p.name for p in meeting.participants) if lang == "ar" else ", ".join(p.name for p in meeting.participants)
    details = [
        [L["date"], meeting.started_at.strftime("%Y-%m-%d") if meeting.started_at else "—"],
        [L["time"], f"{_hm(meeting.started_at)} – {_hm(meeting.ended_at)}"],
        [L["duration"], f"{minutes} {L['minutesUnit']}" if minutes is not None else "—"],
        [L["attendees"], attendees],
    ]
    if chair:
        details.append([L["chair"], chair])
    w.table(None, details)

    if meeting.agenda:
        w.heading(L["agenda"])
        said = {u.agenda_index for u in meeting.utterances}
        for i, item in enumerate(meeting.agenda):
            w.para(f"{i + 1}. {item} — {L['covered'] if i in said else L['skipped']}")

    w.heading(L["summary"])
    for line in (report.summary if report else []) or [L["none"]]:
        w.para(f"• {line}")

    w.heading(L["decisions"])
    decisions = report.decisions if report else []
    if decisions:
        for i, d in enumerate(decisions):
            w.para(f"{i + 1}. {d.text}")
    else:
        w.para(L["none"])

    w.heading(L["tasks"])
    tasks = report.tasks if report else []
    if tasks:
        w.table(
            [L["task"], L["owner"], L["due"], L["status"]],
            [[x.text, names.get(x.owner_id, "—"), x.due_date or "—", L["done"] if x.done else L["open"]] for x in tasks],
        )
    else:
        w.para(L["none"])

    if meeting.utterances:
        w.heading(L["transcript"])
        for u in meeting.utterances:
            on_screen = f"  ({L['slide']} {u.slide.slide})" if u.slide else ""
            w.para(f"[{_hm(u.started_at)}] {names.get(u.participant_id, '?')}: {u.text}{on_screen}", size=9.5)

    # Footer: product name, quietly
    footer = w.doc.sections[0].footer.paragraphs[0]
    if w.rtl:
        _rtl_paragraph(footer)
    _style_run(footer.add_run("LUBB AI"), w.rtl, 8, color=muted)

    out = BytesIO()
    w.doc.save(out)
    return out.getvalue()
